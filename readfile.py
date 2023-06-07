import docx
from io import StringIO
from io import open
from openpyxl import load_workbook
from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfpage import PDFPage

def  readxlsx(path):
  workbook_object = load_workbook(filename=path)
  shett_object = workbook_object.worksheets[0]
  datagroup=[]
  for row in shett_object.rows:
    for col in row:
      if col.value:
        datagroup.append(str(col.value))
  print(datagroup)
  return datagroup

def  readdocx(path):
  doc = docx.Document(path)
  datagroup=[]
  for para in doc.paragraphs:
    if para.text:
      datagroup.append(str(para.text))
  content=' '.join(datagroup)
  return content


def process_pdf(rsrcmgr, device, fp, pagenos=None, maxpages=0, password='', caching=True, check_extractable=True):
  interpreter = PDFPageInterpreter(rsrcmgr, device)
  for page in PDFPage.get_pages(fp, pagenos, maxpages=maxpages, password=password,
              caching=caching, check_extractable=check_extractable):
    interpreter.process_page(page)
  return


def  readpdf(path):
  rsrcmgr = PDFResourceManager()
  retstr = StringIO()
  laparams = LAParams()
  device = TextConverter(rsrcmgr, retstr, laparams=laparams)
  process_pdf(rsrcmgr, device, open(path, "rb"))
  device.close()
  content = retstr.getvalue()
  retstr.close()
  lines = str(content).split("\n")
  return lines



def  readtxt(path):
  with  open(path,encoding='utf8') as f:
    content=f.read()
    f.close()
    return content


def  readfile(path):
   if str(path).split('.')[1].lower() in [ 'docx']:
    return readdocx(path)

   elif str(path).split('.')[1].lower() in ['xlsx']:
    return ' '.join(readxlsx(path))

   elif str(path).split('.')[1].lower() in ['pdf']:
     return ''.join(readpdf(path))
   elif str(path).split('.')[1].lower() in ['txt']:
     return readtxt(path)

